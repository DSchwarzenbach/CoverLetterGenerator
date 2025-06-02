from google import genai
import win32com.client
from docx import Document
from docx2pdf import convert
from pydantic import BaseModel
import streamlit as st
import pythoncom
import requests
from bs4 import BeautifulSoup
import os
import json

#To-do
#Build a good prompt that tells the LLM what stuff to put in what boxes, structure, format etc, might have to change template and variables to make it better
#Add checks for missing information. ie tell them LLM to output null if it cant find a certain feild, then check all fields for potential nulls
#if stuff is null then dont create a new jawn and instead output an error.
#test with option to provide a link, and tell user they need to include all nesecary information somewhere


Output_folder = "CoverLetters"
os.makedirs(Output_folder, exist_ok = True)

def fill_words_template(template_path, output_path, replacements): 
    # 1. Use python-docx to open template and replace placeholders
    try:
        doc = Document(template_path)

        for paragraph in doc.paragraphs:
            for key, value in replacements.items():
                placeholder = f"{{{key}}}"
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, value)

        # Save modified docx
        docx_path = output_path.replace('.pdf', '.docx')
        doc.save(docx_path)
        st.info(f".docx saved to: {docx_path} (will convert to PDF next)")

    except Exception as e:
        st.error(f"Error replacing placeholders in DOCX: {e}")
        return None

    # 2. Use Word COM to open the new DOCX and save as PDF
    pythoncom.CoInitialize()
    word = None
    try:
        abs_docx_path = os.path.abspath(docx_path)
        abs_pdf_path = os.path.abspath(output_path)

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(abs_docx_path)

        # 17 = wdFormatPDF
        doc.SaveAs(abs_pdf_path, FileFormat=17)
        doc.Close(False)

        st.success(f"PDF saved: {abs_pdf_path}")
        return abs_pdf_path

    except Exception as e:
        st.error(f"Error converting DOCX to PDF: {e}")
        st.warning(f"You can manually convert using the DOCX file at {docx_path}")
        return docx_path

    finally:
        if word:
            word.Quit()
        pythoncom.CoUninitialize()

def generate_output(big_blob_of_text, position):
    client = genai.Client(api_key="AIzaSyAqTHAl_xOIIVKf_LVAJVW6LAVRrFnY4Ms")
    #make a thing that puts it into my email template and converts it to a pdf
    #build the prompt using variables and stuff, also maybe do some cooking for a frontend:

    class job_details(BaseModel):
        Hiring_Manager: str
        Body: str
        Company_Name: str
        Company_Address: str
        City_State_Zip: str

    with open('prompt.txt','r') as file:
        text = " ".join(line.rstrip() for line in file)

    text = text.replace("{big_blob_of_text}", big_blob_of_text)

    print(text)

    #Using varibles we can build a pretty structured prompt - general instructions + context about me as an application/
    #what I want it to focus on for each paragraph and then -> then a cover letter body template , and then finally job description

    response = client.models.generate_content(
        model="gemini-2.0-flash", 
        contents= text,
        config={
            "response_mime_type": "application/json",
            "response_schema": job_details,
        },
    )

    parsed_response = json.loads(response.text)
    return parsed_response


st.title("Template-Based Cover Letter Generator")

position = st.text_input("Position")
job_description = st.text_area("Job Description")
job_description_link = st.text_input("Link to job description")

if st.button("Generate Cover Letter"):
    if  job_description:
        content = generate_output(job_description, position)
        template_path = "CoverLetterTemplate.docx"
        output_path = f"{Output_folder}/Schwarzenbach_CoverLetter_{content['Company_Name']}.pdf"
        fill_words_template(template_path, output_path, content)
    else: 
        st.error("Please Fill in Position and Job Desccription")